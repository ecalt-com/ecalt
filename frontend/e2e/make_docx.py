"""Build a small, real .docx for the course-creation E2E (uses python-docx from
the backend venv). Usage: python e2e/make_docx.py <out_path>"""
import sys
from docx import Document

OUT = sys.argv[1] if len(sys.argv) > 1 else "volcanoes.docx"

doc = Document()
doc.add_heading("Volcanoes", level=0)

doc.add_heading("What is a volcano?", level=1)
doc.add_paragraph(
    "A volcano is an opening in the Earth's crust through which molten rock, "
    "ash, and gases escape from below the surface. The molten rock is called "
    "magma while it is underground, and lava once it reaches the surface. "
    "Most volcanoes form near the boundaries of tectonic plates."
)

doc.add_heading("How volcanoes form", level=1)
doc.add_paragraph(
    "Deep inside the Earth, heat and pressure melt rock into magma. Because "
    "magma is less dense than the solid rock around it, it rises through "
    "cracks toward the surface. When the pressure becomes great enough, the "
    "magma erupts through a vent, building up a cone of cooled lava and ash "
    "over many eruptions."
)

doc.add_heading("Types of volcanoes", level=1)
doc.add_paragraph(
    "Shield volcanoes have broad, gentle slopes built from runny lava. "
    "Stratovolcanoes (composite volcanoes) are steep and cone-shaped, made of "
    "alternating layers of lava and ash. Cinder cones are small, steep hills "
    "formed from erupted fragments."
)

doc.add_heading("Why eruptions matter", level=1)
doc.add_paragraph(
    "Eruptions can be destructive, burying towns in ash and lava, but they "
    "also create new land and fertile soil. Scientists called volcanologists "
    "monitor volcanoes to warn people before an eruption, watching for "
    "earthquakes, ground swelling, and escaping gases."
)

doc.save(OUT)
print(f"wrote {OUT}")
