"""End-to-end document extraction job (M3 SC-303/304).

Drives the real `extract` handler through the job queue: a document's bytes are
served by the in-memory storage fake, the handler spawns the sandboxed worker
subprocess to parse them, and chunks land in `document_chunks` with the document
moved to its final status. reportlab (dev-only) builds the PDF fixtures.
"""
import io

import pytest

from app.services import school_ingest_service  # noqa: F401 - registers the 'extract' handler
from app.services import school_jobs as sj
from tests.factories import schools as f

pytestmark = pytest.mark.schools_db

reportlab = pytest.importorskip("reportlab")


def _pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for text in pages:
        c.drawString(72, 750, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def _docx(paragraphs: list[str]) -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("Unit 1", level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _run_extract(db, fake_storage, school, data: bytes, filename="notes.pdf"):
    doc = f.make_document(db, school, filename=filename, status="uploaded")
    fake_storage.put(doc["storage_path"], data)
    sj.enqueue(school["id"], "extract", {"document_id": doc["id"]})
    assert sj.process_one("w1") is True
    row = f.fetch(db, "SELECT status, page_count, error_code FROM source_documents WHERE id = %s", (doc["id"],))[0]
    chunks = f.fetch(db, "SELECT ordinal, text, char_count, heading_path FROM document_chunks "
                         "WHERE document_id = %s ORDER BY ordinal", (doc["id"],))
    return doc, row, chunks


@pytest.fixture
def school(db):
    return f.make_school(db, name="Ingest School")


def test_extract_pdf_to_chunks(school, db, fake_storage):
    data = _pdf(["Photosynthesis lets plants turn light into sugar.",
                 "Chlorophyll in the leaves captures sunlight."])
    doc, row, chunks = _run_extract(db, fake_storage, school, data)
    assert row["status"] == "ready" and row["page_count"] == 2
    assert chunks, "expected document chunks"
    assert "Photosynthesis" in " ".join(c["text"] for c in chunks)
    assert chunks[0]["char_count"] == len(chunks[0]["text"])


def test_extract_docx_to_chunks(school, db, fake_storage):
    data = _docx(["Water flows downhill.", "Rivers carry sediment to the sea."])
    _, row, chunks = _run_extract(db, fake_storage, school, data, filename="lesson.docx")
    assert row["status"] == "ready"
    joined = " ".join(c["text"] for c in chunks)
    assert "Rivers carry sediment" in joined
    # The Heading 1 becomes the chunk's heading_path, not body text.
    assert chunks[0]["heading_path"] == ["Unit 1"]


def test_extract_scanned_pdf_needs_ocr(school, db, fake_storage):
    data = _pdf([" ", " ", " ", " "])  # no text layer → scanned
    _, row, chunks = _run_extract(db, fake_storage, school, data)
    assert row["status"] == "needs_ocr" and row["error_code"] in ("scanned", "no_text")
    assert chunks == []


def test_extract_is_idempotent_on_retry(school, db, fake_storage):
    data = _pdf(["Repeatable extraction run."])
    doc, row, chunks = _run_extract(db, fake_storage, school, data)
    assert row["status"] == "ready"
    first = len(chunks)
    # Re-run the same extract: chunks are replaced, not duplicated.
    sj.enqueue(school["id"], "extract", {"document_id": doc["id"]})
    sj.process_one("w1")
    again = f.fetch(db, "SELECT count(*) AS n FROM document_chunks WHERE document_id = %s", (doc["id"],))
    assert again[0]["n"] == first


def test_extract_missing_document_is_skipped(school, db):
    import uuid
    sj.enqueue(school["id"], "extract", {"document_id": str(uuid.uuid4())})
    sj.process_one("w1")  # handler returns skipped, job succeeds (no crash)
    jobs = f.fetch(db, "SELECT status, result FROM school_jobs WHERE school_id = %s", (school["id"],))
    assert jobs[0]["status"] == "succeeded" and jobs[0]["result"]["status"] == "skipped"
