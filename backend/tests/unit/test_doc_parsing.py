"""Pure parsing + chunking logic for document ingestion (M3 SC-303/304).

These exercise app/services/doc_parsing.py directly on PDFs built with reportlab
and Word files built with python-docx — no subprocess, no DB. reportlab is a
dev-only dependency, so the PDF tests skip where it isn't installed.
"""
import io

import pytest

from app.services import doc_parsing as dp

pytof = pytest.importorskip


def _pdf(pages: list[str]) -> bytes:
    pytof("reportlab")
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for text in pages:
        y = 750
        for line in text.split("\n"):
            c.drawString(72, y, line)
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


def _docx(items: list) -> bytes:
    from docx import Document
    d = Document()
    for kind, content in items:
        if kind == "h1":
            d.add_heading(content, level=1)
        elif kind == "h2":
            d.add_heading(content, level=2)
        elif kind == "body":
            d.add_paragraph(content)
        elif kind == "table":
            t = d.add_table(rows=len(content), cols=len(content[0]))
            for i, row in enumerate(content):
                for j, cell in enumerate(row):
                    t.rows[i].cells[j].text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────

def test_parse_pdf_multipage():
    data = _pdf(["Photosynthesis is how plants make food.", "Chlorophyll absorbs sunlight."])
    out = dp.parse(data, dp.PDF_MIME, max_pages=300)
    assert out["ok"] and out["page_count"] == 2
    assert [b["page_start"] for b in out["blocks"]] == [1, 2]
    assert "Photosynthesis" in out["blocks"][0]["text"]
    assert out["quality"]["scanned"] is False and out["quality"]["garbled"] is False


def test_parse_pdf_scanned_pages_need_ocr():
    # Mostly blank pages (a scan with no text layer) → flagged scanned.
    data = _pdf([" ", " ", " ", "A tiny bit of text here somewhere on the page."])
    out = dp.parse(data, dp.PDF_MIME, max_pages=300)
    assert out["ok"] and out["quality"]["scanned"] is True


def test_parse_pdf_too_many_pages():
    data = _pdf(["one", "two", "three"])
    out = dp.parse(data, dp.PDF_MIME, max_pages=2)
    assert out == {"ok": False, "error_code": "too_many_pages"}


def test_parse_pdf_garbage_is_unreadable():
    assert dp.parse(b"%PDF-not-really", dp.PDF_MIME, max_pages=10)["error_code"] in ("unreadable", "empty")


# ── DOCX ─────────────────────────────────────────────────────────────────────

def test_parse_docx_headings_build_path_and_tables_become_markdown():
    data = _docx([
        ("h1", "Chapter 1"),
        ("body", "Intro paragraph."),
        ("h2", "Section A"),
        ("body", "Section body."),
        ("table", [["Planet", "Moons"], ["Earth", "1"]]),
    ])
    out = dp.parse(data, dp.DOCX_MIME, max_pages=300)
    assert out["ok"]
    by_text = {b["text"]: b["heading_path"] for b in out["blocks"]}
    assert by_text["Intro paragraph."] == ["Chapter 1"]
    assert by_text["Section body."] == ["Chapter 1", "Section A"]
    table = next(b["text"] for b in out["blocks"] if b["text"].startswith("| Planet"))
    assert "| --- | --- |" in table and "| Earth | 1 |" in table


def test_parse_docx_rejects_macros():
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("word/document.xml", "<w:document/>")
        zf.writestr("word/vbaProject.bin", b"\x00macro")
    assert dp.parse(buf.getvalue(), dp.DOCX_MIME, max_pages=10) == {"ok": False, "error_code": "macro"}


# ── Quality + chunking ───────────────────────────────────────────────────────

def test_sensible_ratio_flags_broken_encoding():
    assert dp._sensible_ratio("Plants need sunlight, water, and CO2.") > 0.9
    pua = "".join(chr(0xE000 + (i % 200)) for i in range(300))  # private-use glyphs
    assert dp._sensible_ratio(pua) < 0.6


def test_chunk_blocks_splits_on_section_and_size():
    blocks = [
        {"text": "a" * 100, "heading_path": ["Intro"], "page_start": 1, "page_end": 1},
        {"text": "b" * 100, "heading_path": ["Intro"], "page_start": 1, "page_end": 1},
        {"text": "c" * 100, "heading_path": ["Body"], "page_start": 2, "page_end": 2},
    ]
    chunks = dp.chunk_blocks(blocks, max_chars=400_000, target_chars=1000)
    # Intro's two blocks merge; the heading change starts a new chunk.
    assert len(chunks) == 2
    assert chunks[0]["heading_path"] == ["Intro"] and chunks[0]["page_start"] == 1
    assert chunks[1]["heading_path"] == ["Body"] and chunks[1]["page_start"] == 2
    assert [c["ordinal"] for c in chunks] == [0, 1]
    assert chunks[0]["char_count"] == len(chunks[0]["text"])


def test_chunk_blocks_respects_target_size():
    blocks = [{"text": "x" * 500, "heading_path": [], "page_start": None, "page_end": None} for _ in range(5)]
    chunks = dp.chunk_blocks(blocks, max_chars=400_000, target_chars=1000)
    assert len(chunks) >= 2 and all(c["char_count"] <= 1600 for c in chunks)


def test_chunk_blocks_strips_repeated_headers():
    pages = [{"text": f"ACME SCHOOL WORKSHEET\nReal content on page {i}.", "heading_path": [],
              "page_start": i, "page_end": i} for i in range(1, 7)]
    chunks = dp.chunk_blocks(pages, max_chars=400_000, target_chars=100_000)
    joined = " ".join(c["text"] for c in chunks)
    assert "ACME SCHOOL WORKSHEET" not in joined and "Real content on page 1." in joined


def test_chunk_blocks_caps_total_chars():
    blocks = [{"text": "z" * 1000, "heading_path": [], "page_start": None, "page_end": None} for _ in range(50)]
    chunks = dp.chunk_blocks(blocks, max_chars=5000, target_chars=2000)
    assert sum(c["char_count"] for c in chunks) <= 6000  # stops near the cap
